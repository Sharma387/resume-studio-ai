"""Phase-12 smoke: real artifacts from a real stored resume.

Generates PDF / DOCX / HTML through the canonical export API for a matrix of
registered layouts and themes and validates the actual bytes on disk — never
mocked renderers or stubbed content.
"""

import uuid
from io import BytesIO
from pathlib import Path

import pytest
from httpx import ASGITransport, AsyncClient
from pypdf import PdfReader

from app.main import app
from app.models.resume import Resume
from app.services.storage_service import save_resume

REPRESENTATIVE = (
    ("executive", "blue"),
    ("sidebar", "gold"),
    ("classic", "slate"),
)

EXPECTED_CONTENT = {
    "name": "Regina Artifact",
    "title": "Principal Platforms Engineer",
    "experience": "Northwind Systems",
    "education": "University of Auckland",
    "skills": "Terraform",
}


@pytest.fixture
def client() -> AsyncClient:
    return AsyncClient(transport=ASGITransport(app=app), base_url="http://test")


def _save_real_resume() -> str:
    resume_id = uuid.uuid4().hex
    resume = Resume(
        user_id="dev-user",
        full_name=EXPECTED_CONTENT["name"],
        email="regina@example.com",
        phone="+64 21 555 0199",
        location="Auckland, NZ",
        professional_title=EXPECTED_CONTENT["title"],
        summary="Platform leader delivering resilient infrastructure at scale.",
        experience=[
            {
                "company": "Northwind Systems",
                "title": "Principal Engineer",
                "location": "Auckland, NZ",
                "start_date": "2019-03",
                "current": True,
                "description": ["Owned the multi-region platform", "Cut incident MTTR by 60%"],
            },
            {
                "company": "Contoso Labs",
                "title": "Staff Engineer",
                "start_date": "2015-06",
                "end_date": "2019-02",
                "description": ["Built the streaming data pipeline"],
            },
        ],
        education=[
            {
                "institution": "University of Auckland",
                "degree": "Master of Engineering",
                "field": "Software",
                "gpa": 3.8,
            }
        ],
        skills=[{"category": "IaC", "skills": ["Terraform", "Pulumi"]}],
        certifications=[{"name": "AWS Solutions Architect Professional", "issuer": "AWS", "date": "2021"}],
        projects=[{"name": "Atlas", "description": "Global state plane", "technologies": ["Go"]}],
        languages=[{"name": "English", "proficiency": "Native"}],
    )
    save_resume(resume_id, resume)
    return resume_id


@pytest.mark.asyncio
async def test_export_pdf_artifacts_valid(client, tmp_path: Path):
    resume_id = _save_real_resume()
    for layout_id, theme_id in REPRESENTATIVE:
        response = await client.post(
            f"/api/v1/resume/{resume_id}/export",
            json={"layout_id": layout_id, "theme_id": theme_id, "format": "pdf"},
        )
        assert response.status_code == 200, (layout_id, theme_id)
        path = tmp_path / f"{layout_id}-{theme_id}.pdf"
        path.write_bytes(response.content)
        assert path.read_bytes()[:5] == b"%PDF-"
        reader = PdfReader(str(path))
        assert len(reader.pages) > 0
        text = " ".join((page.extract_text() or "") for page in reader.pages)
        for needle in EXPECTED_CONTENT.values():
            assert needle in text, (layout_id, theme_id, needle)


@pytest.mark.asyncio
async def test_export_docx_artifacts_valid(client, tmp_path: Path):
    from docx import Document

    resume_id = _save_real_resume()
    for layout_id, theme_id in REPRESENTATIVE:
        response = await client.post(
            f"/api/v1/resume/{resume_id}/export",
            json={"layout_id": layout_id, "theme_id": theme_id, "format": "docx"},
        )
        assert response.status_code == 200, (layout_id, theme_id)
        path = tmp_path / f"{layout_id}-{theme_id}.docx"
        path.write_bytes(response.content)
        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        text = "\n".join(parts)
        assert doc.paragraphs or doc.tables
        assert EXPECTED_CONTENT["name"] in text
        assert EXPECTED_CONTENT["experience"] in text
        assert EXPECTED_CONTENT["education"] in text
        assert EXPECTED_CONTENT["skills"] in text


@pytest.mark.asyncio
async def test_export_html_artifacts_valid(client, tmp_path: Path):
    resume_id = _save_real_resume()
    for layout_id, theme_id in REPRESENTATIVE:
        response = await client.post(
            f"/api/v1/resume/{resume_id}/export",
            json={"layout_id": layout_id, "theme_id": theme_id, "format": "html"},
        )
        assert response.status_code == 200, (layout_id, theme_id)
        html = response.content.decode("utf-8")
        assert html.startswith("<!DOCTYPE html>")
        assert 'class="resume-region resume-region-' in html
        assert 'data-region="' in html
        assert EXPECTED_CONTENT["name"] in html
        assert EXPECTED_CONTENT["title"] in html
        assert EXPECTED_CONTENT["experience"] in html
        assert EXPECTED_CONTENT["education"] in html
        assert EXPECTED_CONTENT["skills"] in html
        assert 'id="rsai-theme"' in html


@pytest.mark.asyncio
async def test_cover_letter_pdf_still_valid(client, tmp_path: Path):
    """ReportLab cover-letter PDF generation remains intact end-to-end."""
    from app.models.cover_letter import CoverLetter
    from app.services.cover_letter_pdf import generate_cover_letter_pdf
    from app.services.storage_service import load_resume, save_cover_letter

    resume_id = _save_real_resume()
    letter = CoverLetter(
        user_id="dev-user",
        id="cl-smoke",
        resume_id=resume_id,
        company_name="Northwind Systems",
        hiring_manager="Hiring Manager",
        role_title="Principal Platform Engineer",
        subject="Application for Principal Platform Engineer",
        content="Dear Hiring Manager,\n\nI am excited to apply for the role.\n\nSincerely",
    )
    resume = load_resume(resume_id)
    save_cover_letter(letter)
    path = generate_cover_letter_pdf(resume_id, letter.id, letter, resume)
    assert path.exists()
    blob = path.read_bytes()
    assert blob[:5] == b"%PDF-"
    reader = PdfReader(BytesIO(blob))
    assert len(reader.pages) > 0
    text = " ".join((page.extract_text() or "") for page in reader.pages)
    assert EXPECTED_CONTENT["name"] in text
    assert "Northwind Systems" in text
    assert "Principal Platform Engineer" in text
