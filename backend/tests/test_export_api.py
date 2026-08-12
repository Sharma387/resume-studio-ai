"""Tests for the unified RenderTree export API (Layout Engine)."""

import re
import uuid
from collections import Counter

import pytest
from httpx import ASGITransport, AsyncClient
from pypdf import PdfReader

from app.main import app
from app.models.resume import Resume
from app.rendering.context import OutputFormat
from app.rendering.export_service import ExportFormat, ExportService
from app.services.storage_service import save_resume


@pytest.fixture
def client() -> AsyncClient:
    return AsyncClient(transport=ASGITransport(app=app), base_url="http://test")


def _resume(user_id="test") -> Resume:
    return Resume(
        user_id=user_id,
        full_name="Sharma Rajasekar",
        email="sharma@test.com",
        phone="+1-555-0100",
        location="Auckland, New Zealand",
        linkedin="https://linkedin.com/in/sharma",
        website="https://sharma.dev",
        professional_title="Senior Transformation & Infrastructure Leader",
        summary="Transformation leader focused on distributed systems and platform reliability.",
        experience=[
            {"company": "Acme Corporation International", "title": "Senior Software Engineering Lead",
             "location": "San Francisco, CA", "start_date": "2021", "current": True,
             "description": ["Led a 12-engineer platform organization", "Designed multi-region streaming"]},
            {"company": "Beta Inc", "title": "Transformation Programme Manager",
             "start_date": "2018", "end_date": "2021", "description": ["Built a distributed scheduler"]},
        ],
        education=[{"institution": "Carnegie Mellon University", "degree": "Master of Science", "field": "Computer Science", "gpa": 3.9}],
        skills=[{"category": "Languages", "skills": ["Python", "Go", "Rust"]}],
        certifications=[{"name": "AWS Solutions Architect", "issuer": "Amazon Web Services", "date": "2022"}],
        projects=[{"name": "Project Alpha", "description": "Distributed scheduler", "url": "https://github.com/example/alpha", "technologies": ["Go", "PostgreSQL"]}],
        awards=[{"name": "Transformation Excellence Award", "issuer": "Acme Corporation", "date": "2023-06"}],
        languages=[{"name": "English", "proficiency": "Native"}, {"name": "French", "proficiency": "Professional"}],
    )


def _save_resume(user_id="test") -> str:
    resume_id = uuid.uuid4().hex
    save_resume(resume_id, _resume(user_id))
    return resume_id


def _body_text(html: str) -> str:
    stripped = re.sub(r"<style.*?</style>", "", html, flags=re.S)
    return " ".join(re.sub(r"<[^>]+>", " ", stripped).split())


def _pdf_text(data: bytes) -> str:
    return " ".join((page.extract_text() or "") for page in PdfReader(io_bytes(data)).pages)


def _docx_text(data: bytes) -> str:
    from docx import Document

    doc = Document(io_bytes(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def io_bytes(data: bytes):
    from io import BytesIO

    return BytesIO(data)


def _tokens(text: str) -> Counter:
    return Counter(text.split())


async def _export(client: AsyncClient, resume_id: str, body: dict) -> tuple[int, dict, bytes]:
    response = await client.post(f"/api/v1/resume/{resume_id}/export", json=body)
    return response.status_code, dict(response.headers), response.content


# ── basic exports ─────────────────────────────────────────────────────────────


class TestExportEndpoint:
    async def test_pdf_export(self, client):
        rid = _save_resume()
        status, headers, content = await _export(client, rid, {"layout_id": "sidebar", "theme_id": "blue", "format": "pdf"})
        assert status == 200
        assert headers["content-type"] == "application/pdf"
        assert content[:5] == b"%PDF-"
        assert "Sharma Rajasekar" in _pdf_text(content)
        assert 'filename="resume-sidebar.pdf"' in headers["content-disposition"]

    async def test_docx_export(self, client):
        rid = _save_resume()
        status, headers, content = await _export(client, rid, {"layout_id": "executive", "theme_id": "blue", "format": "docx"})
        assert status == 200
        assert headers["content-type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        assert content[:2] == b"PK"
        assert "Sharma Rajasekar" in _docx_text(content)

    async def test_html_export(self, client):
        rid = _save_resume()
        status, headers, content = await _export(client, rid, {"layout_id": "classic", "theme_id": "blue", "format": "html"})
        assert status == 200
        assert headers["content-type"].startswith("text/html")
        assert "Sharma Rajasekar" in content.decode("utf-8")


# ── validation / errors ───────────────────────────────────────────────────────


class TestValidation:
    async def test_unknown_layout_returns_404(self, client):
        rid = _save_resume()
        status, _, _ = await _export(client, rid, {"layout_id": "nope", "theme_id": "blue", "format": "pdf"})
        assert status == 404

    async def test_unknown_theme_returns_404(self, client):
        rid = _save_resume()
        status, _, _ = await _export(client, rid, {"layout_id": "sidebar", "theme_id": "nope", "format": "pdf"})
        assert status == 404

    async def test_unsupported_format_returns_422(self, client):
        rid = _save_resume()
        status, _, _ = await _export(client, rid, {"layout_id": "sidebar", "theme_id": "blue", "format": "png"})
        assert status == 422

    async def test_unknown_resume_returns_404(self, client):
        status, _, _ = await _export(client, "nonexistent", {"layout_id": "sidebar", "theme_id": "blue", "format": "pdf"})
        assert status == 404

    async def test_another_users_resume_returns_404(self, client):
        rid = _save_resume(user_id="other")
        status, _, _ = await _export(client, rid, {"layout_id": "sidebar", "theme_id": "blue", "format": "pdf"})
        assert status == 404


# ── same content across formats ───────────────────────────────────────────────


class TestSameContentAcrossFormats:
    async def test_same_values_in_all_formats(self, client):
        rid = _save_resume()
        artifacts = {}
        for fmt in ("html", "pdf", "docx"):
            status, _, content = await _export(client, rid, {"layout_id": "sidebar", "theme_id": "blue", "format": fmt})
            assert status == 200
            artifacts[fmt] = content

        html_text = _body_text(artifacts["html"].decode("utf-8"))
        pdf_text = _pdf_text(artifacts["pdf"])
        docx_text = _docx_text(artifacts["docx"])
        for text in (html_text, pdf_text, docx_text):
            assert "Sharma Rajasekar" in text
            assert "Senior Transformation" in text
            assert "Acme Corporation International" in text
            assert "Carnegie Mellon University" in text
            assert "Python" in text
            assert "AWS Solutions Architect" in text
            assert "Project Alpha" in text
            assert "Transformation Excellence Award" in text
            assert "English" in text and "French" in text
            assert "sharma@test.com" in text


# ── service dispatch: same RenderTree to all renderers ────────────────────────


class TestServiceDispatch:
    def test_same_tree_renderer_receives_tree(self, monkeypatch):
        recorded = []

        class FakeRenderer:
            def __init__(self):
                self.trees = []

            def render(self, tree, *, theme=None):
                recorded.append((tree, theme))
                return b"fake"

        from app.rendering import export_service

        monkeypatch.setitem(export_service._RENDERERS, OutputFormat.HTML, FakeRenderer)
        monkeypatch.setitem(export_service._RENDERERS, OutputFormat.PDF, FakeRenderer)
        monkeypatch.setitem(export_service._RENDERERS, OutputFormat.DOCX, FakeRenderer)

        service = ExportService()
        resume = _resume()
        for fmt in (ExportFormat.HTML, ExportFormat.PDF, ExportFormat.DOCX):
            result = service.export(resume, layout_id="sidebar", theme_id="blue", output_format=fmt)
            assert result.content == b"fake"
        assert len(recorded) == 3
        # All three renderers received the same RenderTree (same serialization).
        trees = [tree.model_dump_json() for tree, _ in recorded]
        assert trees[0] == trees[1] == trees[2]


# ── theme separation via API ──────────────────────────────────────────────────


class TestThemeViaAPI:
    async def test_blue_vs_gold_content_identical(self, client):
        rid = _save_resume()
        _, _, blue = await _export(client, rid, {"layout_id": "sidebar", "theme_id": "blue", "format": "pdf"})
        _, _, gold = await _export(client, rid, {"layout_id": "sidebar", "theme_id": "gold", "format": "pdf"})
        assert _tokens(_pdf_text(blue)) == _tokens(_pdf_text(gold))
        assert "Sharma Rajasekar" in _pdf_text(blue) and "Sharma Rajasekar" in _pdf_text(gold)
