"""Production hardening: every registered layout previews and exports cleanly.

The retirement phase verified executive / sidebar / modern / classic. This test
enumerates the ACTUAL registry (``/resume/layouts``) and drives every registered
layout through preview + PDF/DOCX/HTML export, validating the resulting
artifacts and verifying layout-specific structure matches the LayoutDefinition
declared in the reference registry (source of truth — no invented
expectations).
"""

import re
import uuid
from collections import Counter
from html.parser import HTMLParser
from io import BytesIO

import pytest
from httpx import ASGITransport, AsyncClient
from pypdf import PdfReader

from app.main import app
from app.models.resume import Certification as ResumeCertification
from app.models.resume import Education as ResumeEducation
from app.models.resume import Experience as ResumeExperience
from app.models.resume import Resume
from app.models.resume import Skill as ResumeSkill
from app.rendering.layout.reference_layouts import REFERENCE_LAYOUTS
from app.services.storage_service import save_resume

#: Every layout the reference registry actually registers. This guard fails if a
#: newly registered layout is ever left untested.
REGISTERED_LAYOUT_IDS = tuple(layout.layout_id for layout in REFERENCE_LAYOUTS)

#: Where each registered layout places the canonical ``skills`` section, derived
#: from placement rules / region ``allowed_sections`` declarations (source of
#: truth in the reference layout definitions).
EXPECTED_SKILLS_REGION = {
    "executive": "main",
    "modern": "secondary",
    "sidebar": "sidebar",
    "timeline": "main",
    "classic": "secondary",
    "minimal": "main",
}


class _StructureParser(HTMLParser):
    _VOID = {"hr", "img", "br", "meta", "link", "input"}

    def __init__(self) -> None:
        super().__init__()
        self._stack: list[tuple[str, str | None]] = []
        self.section_regions: dict[str, str | None] = {}
        self.regions: list[str] = []

    def handle_starttag(self, tag: str, attrs) -> None:
        if tag in self._VOID:
            return
        attrs = dict(attrs)
        region = attrs.get("data-region") if tag == "div" else None
        self._stack.append((tag, region))
        if tag == "div" and region:
            self.regions.append(region)
        if tag == "section":
            enclosing: str | None = None
            for _, region_id in reversed(self._stack[:-1]):
                if region_id:
                    enclosing = region_id
                    break
            self.section_regions[attrs.get("data-section")] = enclosing

    def handle_endtag(self, tag: str) -> None:
        if tag in self._VOID:
            return
        for index in range(len(self._stack) - 1, -1, -1):
            if self._stack[index][0] == tag:
                del self._stack[index:]
                return


def _structure(html: str) -> _StructureParser:
    parser = _StructureParser()
    parser.feed(html)
    parser.close()
    return parser


def _body_text(html: str) -> str:
    without_style = re.sub(r"<style.*?</style>", "", html, flags=re.S)
    stripped = re.sub(r"<[^>]+>", " ", without_style)
    return " ".join(stripped.split())


def _tokens(text: str) -> Counter:
    return Counter(text.split())


def _definition(layout_id: str):
    for layout in REFERENCE_LAYOUTS:
        if layout.layout_id == layout_id:
            return layout
    raise AssertionError(f"layout {layout_id!r} not registered")


def _pdf_text(data: bytes) -> str:
    return " ".join((page.extract_text() or "") for page in PdfReader(BytesIO(data)).pages)


def _docx_joined(data: bytes) -> str:
    from docx import Document

    doc = Document(BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def _docx_table_count(data: bytes) -> int:
    from docx import Document

    return len(Document(BytesIO(data)).tables)


@pytest.fixture
def client() -> AsyncClient:
    return AsyncClient(transport=ASGITransport(app=app), base_url="http://test")


def _save_resume() -> str:
    resume_id = uuid.uuid4().hex
    resume = Resume(
        user_id="dev-user",
        full_name="Jane Doe",
        email="jane@test.com",
        professional_title="Senior Software Engineer",
        summary="Full-stack engineer with 8 years building platforms.",
        experience=[
            ResumeExperience(company="Acme", title="Senior Engineer", start_date="2016", current=True),
            ResumeExperience(company="Beta Inc", title="Engineer", start_date="2014", end_date="2016"),
        ],
        education=[ResumeEducation(institution="MIT", degree="B.Sc.", field="Computer Science")],
        skills=[ResumeSkill(category="Languages", skills=["Python", "Go"])],
        certifications=[ResumeCertification(name="AWS Certified", issuer="Amazon")],
    )
    save_resume(resume_id, resume)
    return resume_id


async def _preview(client: AsyncClient, resume_id: str, **params) -> tuple[dict, str]:
    url = f"/api/v1/resume/{resume_id}/preview"
    if params:
        url += "?" + "&".join(f"{key}={value}" for key, value in params.items())
    response = await client.get(url)
    assert response.status_code == 200, (params, response.json())
    body = response.json()
    file_response = await client.get(body["data"]["preview_url"])
    assert file_response.status_code == 200
    return body, file_response.text


async def _export(client: AsyncClient, resume_id: str, body: dict) -> tuple[int, dict, bytes]:
    response = await client.post(f"/api/v1/resume/{resume_id}/export", json=body)
    return response.status_code, dict(response.headers), response.content


def _verify_artifact(fmt: str, content: bytes, headers: dict) -> None:
    if fmt == "pdf":
        assert headers["content-type"] == "application/pdf"
        assert content[:5] == b"%PDF-"
    elif fmt == "docx":
        assert content[:2] == b"PK"
        assert "wordprocessingml" in headers["content-type"]
    elif fmt == "html":
        assert headers["content-type"].startswith("text/html")
        assert b"<!DOCTYPE html>" in content
        assert b"Jane Doe" in content


@pytest.mark.asyncio
async def test_registry_lists_all_reference_layouts(client):
    response = await client.get("/api/v1/resume/layouts")
    assert response.status_code == 200
    ids = {item["layout_id"] for item in response.json()["data"]}
    assert ids == set(REGISTERED_LAYOUT_IDS)
    assert len(ids) >= 4


@pytest.mark.asyncio
async def test_every_registered_layout_previews_and_exports(client):
    resume_id = _save_resume()
    for layout_id in REGISTERED_LAYOUT_IDS:
        # Preview renders valid HTML containing every declared region.
        _, html = await _preview(client, resume_id, layout_id=layout_id, theme="blue")
        structure = _structure(html)
        definition = _definition(layout_id)
        for region in definition.regions:
            assert region.identifier in structure.regions, f"{layout_id} missing {region.identifier}"
        assert "Jane Doe" in _body_text(html)
        assert "Senior Engineer" in _body_text(html)
        assert "Acme" in _body_text(html)

        # Every format produces a valid artifact with full resume content.
        for fmt in ("pdf", "docx", "html"):
            status, headers, content = await _export(
                client, resume_id, {"layout_id": layout_id, "theme_id": "blue", "format": fmt}
            )
            assert status == 200, (layout_id, fmt)
            _verify_artifact(fmt, content, headers)

        # PDF: valid signature, >0 pages, core sections present.
        status, _, pdf = await _export(client, resume_id, {"layout_id": layout_id, "theme_id": "blue", "format": "pdf"})
        assert status == 200
        reader = PdfReader(BytesIO(pdf))
        assert len(reader.pages) > 0
        text = _pdf_text(pdf)
        assert "Jane Doe" in text
        assert "Acme" in text
        assert "MIT" in text
        assert "Python" in text

        # DOCX: opens, content present.
        status, _, docx = await _export(
            client, resume_id, {"layout_id": layout_id, "theme_id": "blue", "format": "docx"}
        )
        assert status == 200
        joined = _docx_joined(docx)
        assert "Jane Doe" in joined
        assert "Acme" in joined
        assert "MIT" in joined
        assert "Python" in joined


@pytest.mark.asyncio
async def test_layout_structure_differs_per_registry_definition(client):
    resume_id = _save_resume()
    bodies: dict[str, str] = {}
    structures: dict[str, _StructureParser] = {}
    for layout_id in REGISTERED_LAYOUT_IDS:
        _, html = await _preview(client, resume_id, layout_id=layout_id, theme="blue")
        bodies[layout_id] = _body_text(html)
        structures[layout_id] = _structure(html)

    # Skills placement matches the declared placement for every layout.
    for layout_id, expected_region in EXPECTED_SKILLS_REGION.items():
        assert structures[layout_id].section_regions.get("skills") == expected_region, layout_id

    # Multi-column layouts render a DOCX table row; single-column layouts don't.
    status, _, sidebar_docx = await _export(
        client, resume_id, {"layout_id": "sidebar", "theme_id": "blue", "format": "docx"}
    )
    assert status == 200
    assert _docx_table_count(sidebar_docx) >= 1
    status, _, classic_docx = await _export(
        client, resume_id, {"layout_id": "classic", "theme_id": "blue", "format": "docx"}
    )
    assert status == 200
    assert _docx_table_count(classic_docx) >= 1
    status, _, minimal_docx = await _export(
        client, resume_id, {"layout_id": "minimal", "theme_id": "blue", "format": "docx"}
    )
    assert status == 200
    assert _docx_table_count(minimal_docx) == 0

    # The same resume yields identical content tokens in every layout.
    counters = [_tokens(bodies[layout_id]) for layout_id in REGISTERED_LAYOUT_IDS]
    assert all(counters[0] == counter for counter in counters[1:])


@pytest.mark.asyncio
async def test_unknown_layout_404_for_every_export_format(client):
    resume_id = _save_resume()
    for fmt in ("pdf", "docx", "html"):
        status, _, _ = await _export(client, resume_id, {"layout_id": "nope", "theme_id": "blue", "format": fmt})
        assert status == 404, fmt
