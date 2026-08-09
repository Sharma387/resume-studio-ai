"""Tests for the RenderTree → DOCX renderer (Layout Engine)."""

import ast
import importlib
import inspect
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document

from app.rendering.content import ContentView, Profile
from app.rendering.content.models import (
    AwardEntry,
    CertificationEntry,
    EducationEntry,
    ExperienceEntry,
    LanguageEntry,
    ProjectEntry,
    SkillGroup,
)
from app.rendering.layout.reference_layouts import classic_layout, executive_layout, modern_layout, sidebar_layout
from app.rendering.layout_html import render_layout_docx
from app.rendering.renderers.tree_docx_renderer import RenderTreeDOCXRenderer
from app.rendering.theme.reference_themes import blue_theme, gold_theme
from app.rendering.tree import NodeKind, RenderNode
from app.rendering.tree.validator import TreeValidationError


def _comprehensive_cvm() -> ContentView:
    return ContentView(
        stable_id="resume.docx",
        profile=Profile(
            full_name="Sharma Rajasekar",
            professional_title="Senior Transformation & Infrastructure Leader",
            email="sharma@test.com",
            phone="+1-555-0100",
            location="Auckland, New Zealand",
            linkedin="https://linkedin.com/in/sharma",
            github="https://github.com/sharma",
            website="https://sharma.dev",
        ),
        summary="Transformation leader focused on distributed systems and platform reliability.",
        experience=(
            ExperienceEntry(company="Acme Corporation International", title="Senior Software Engineering Lead",
                            location="San Francisco, CA", start_date="2021", current=True,
                            description=("Led a 12-engineer platform organization", "Designed multi-region streaming")),
            ExperienceEntry(company="Beta Inc", title="Transformation Programme Manager",
                            start_date="2018", end_date="2021", description=("Built a distributed scheduler",)),
            ExperienceEntry(company="Gamma Labs", title="Senior Software Engineer",
                            start_date="2016", end_date="2018", description=("Developed observability tooling",)),
        ),
        education=(
            EducationEntry(institution="Carnegie Mellon University", degree="Master of Science", field="Computer Science", gpa=3.9),
            EducationEntry(institution="University of Auckland", degree="Bachelor of Engineering", field="Software", gpa=3.6),
        ),
        skills=(
            SkillGroup(category="Languages", skills=("Python", "Go", "Rust")),
            SkillGroup(category="Cloud", skills=("AWS", "Azure", "Kubernetes")),
        ),
        certifications=(
            CertificationEntry(name="AWS Solutions Architect", issuer="Amazon Web Services", date="2022"),
            CertificationEntry(name="PMP", issuer="PMI", date="2021"),
        ),
        projects=(
            ProjectEntry(name="Project Alpha", description="Distributed scheduler", url="https://github.com/example/alpha", technologies=("Go", "PostgreSQL")),
            ProjectEntry(name="Project Beta", description="Streaming reliability platform", technologies=("Rust", "Kafka")),
            ProjectEntry(name="Project Gamma", description="Observability platform", technologies=("Python", "ClickHouse")),
        ),
        awards=(
            AwardEntry(title="Transformation Excellence Award", issuer="Acme Corporation", date="2023-06"),
            AwardEntry(title="Innovation Leadership Award", issuer="Industry Forum", date="2022-11"),
        ),
        languages=(
            LanguageEntry(name="English", proficiency="Native"),
            LanguageEntry(name="Tamil", proficiency="Native"),
            LanguageEntry(name="French", proficiency="Professional"),
        ),
    )


def _long_cvm() -> ContentView:
    return ContentView(
        stable_id="resume.long",
        profile=Profile(full_name="Alexandra Rivera", professional_title="Distinguished Platform Engineer"),
        summary="Very long professional summary that keeps going across multiple lines of text " * 4,
        experience=tuple(
            ExperienceEntry(
                company=f"Company {i} International Solutions Group",
                title=f"Senior Principal Engineering Manager {i}",
                start_date=f"20{10 + i}",
                end_date=None if i == 5 else f"20{15 + i}",
                current=i == 5,
                description=("Led distributed teams across regions", "Designed streaming platforms", "Reduced costs by 35%"),
            )
            for i in range(6)
        ),
        education=tuple(EducationEntry(institution=f"University {i}", degree="Ph.D.", field="Computer Science") for i in range(3)),
        skills=tuple(SkillGroup(category=f"Group {i}", skills=("Python", "Go", "Rust", "AWS", "Kubernetes", "Kafka", "Docker", "Terraform")) for i in range(2)),
        certifications=tuple(CertificationEntry(name=f"Certification {i}", issuer="Vendor", date="2022") for i in range(4)),
        projects=tuple(ProjectEntry(name=f"Project {i}", description="A long project description", technologies=("Go", "Rust")) for i in range(3)),
        awards=tuple(AwardEntry(title=f"Award {i}", issuer="Org", date="2023") for i in range(2)),
        languages=tuple(LanguageEntry(name=f"Language {i}", proficiency="Native") for i in range(3)),
    )


def _open(data: bytes) -> Document:
    return Document(BytesIO(data))


def _doc_text(doc: Document) -> str:
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


# ── 1/2: valid DOCX ───────────────────────────────────────────────────────────


class TestValidDOCX:
    def test_signature_and_opens(self):
        data = render_layout_docx(_comprehensive_cvm(), executive_layout(), blue_theme())
        assert data[:2] == b"PK"  # ZIP/DOCX container
        doc = _open(data)
        assert len(doc.paragraphs) > 0


# ── substantive content preservation ──────────────────────────────────────────


class TestContentPreservation:
    def test_all_values_present(self):
        doc = _open(render_layout_docx(_comprehensive_cvm(), sidebar_layout(), blue_theme()))
        text = _doc_text(doc)
        for expected in (
            "Sharma Rajasekar",
            "Senior Transformation",
            "Infrastructure Leader",
            "Acme Corporation International",
            "Senior Software Engineering Lead",
            "Transformation Programme Manager",
            "Carnegie Mellon University",
            "Master of Science",
            "AWS",
            "Azure",
            "Kubernetes",
            "AWS Solutions Architect",
            "PMP",
            "Project Alpha",
            "Project Beta",
            "Project Gamma",
            "Transformation Excellence Award",
            "Innovation Leadership Award",
            "English",
            "Tamil",
            "French",
            "Native",
            "Professional",
            "sharma@test.com",
            "Auckland, New Zealand",
        ):
            assert expected in text, expected

    def test_structured_sections(self):
        doc = _open(render_layout_docx(_comprehensive_cvm(), executive_layout(), blue_theme()))
        text = _doc_text(doc)
        assert "EXPERIENCE" in text.upper()
        assert "EDUCATION" in text.upper()
        assert "SKILLS" in text.upper()
        assert "CERTIFICATIONS" in text.upper()
        assert "PROJECTS" in text.upper()
        assert "AWARDS" in text.upper()
        assert "LANGUAGES" in text.upper()

    def test_bullets_rendered(self):
        doc = _open(render_layout_docx(_comprehensive_cvm(), sidebar_layout(), blue_theme()))
        text = _doc_text(doc)
        assert "Led a 12-engineer platform organization" in text
        assert "Python" in text and "Go" in text and "Rust" in text


# ── layout structure ──────────────────────────────────────────────────────────


class TestLayoutStructure:
    def test_single_column_layouts_have_no_table(self):
        cvm = _comprehensive_cvm()
        for layout in (executive_layout(), classic_layout()):
            doc = _open(render_layout_docx(cvm, layout, blue_theme()))
            assert len(doc.tables) == 0, f"{layout.layout_id} should be single-column"

    def test_two_column_layouts_have_table(self):
        cvm = _comprehensive_cvm()
        for layout in (sidebar_layout(), modern_layout()):
            doc = _open(render_layout_docx(cvm, layout, blue_theme()))
            assert len(doc.tables) == 1
            assert len(doc.tables[0].columns) == 2

    def test_sidebar_table_cells_hold_distinct_content(self):
        doc = _open(render_layout_docx(_comprehensive_cvm(), sidebar_layout(), blue_theme()))
        table = doc.tables[0]
        cell0 = table.cell(0, 0).text
        cell1 = table.cell(0, 1).text
        assert "PROFESSIONAL SUMMARY" in cell0.upper() or "SUMMARY" in cell0.upper()
        assert "SKILLS" in cell1.upper()
        assert "AWS" in cell1  # skills/certs in the sidebar cell

    def test_same_content_across_layouts(self):
        cvm = _comprehensive_cvm()
        texts = []
        for layout in (executive_layout(), sidebar_layout(), modern_layout(), classic_layout()):
            doc = _open(render_layout_docx(cvm, layout, blue_theme()))
            texts.append(_doc_text(doc))
        # Same substantive content regardless of layout.
        for text in texts:
            assert "Sharma Rajasekar" in text
            assert "Project Alpha" in text
            assert "Transformation Excellence Award" in text
            assert "English" in text


# ── theme separation ─────────────────────────────────────────────────────────


class TestThemeSeparation:
    def test_theme_changes_style_not_content(self):
        cvm = _comprehensive_cvm()
        blue = _open(render_layout_docx(cvm, sidebar_layout(), blue_theme()))
        gold = _open(render_layout_docx(cvm, sidebar_layout(), gold_theme()))
        assert _doc_text(blue) == _doc_text(gold)
        assert len(blue.tables) == len(gold.tables) == 1

    def test_heading_color_differs_by_theme(self):
        cvm = _comprehensive_cvm()

        def heading_colors(doc: Document):
            colors = set()
            for p in doc.paragraphs:
                if p.text.strip().upper() in ("PROFILE", "SUMMARY", "EXPERIENCE", "EDUCATION"):
                    for run in p.runs:
                        if run.font.color is not None and run.font.color.rgb is not None:
                            colors.add(str(run.font.color.rgb))
            return colors

        blue_colors = heading_colors(_open(render_layout_docx(cvm, sidebar_layout(), blue_theme())))
        gold_colors = heading_colors(_open(render_layout_docx(cvm, sidebar_layout(), gold_theme())))
        assert blue_colors and gold_colors
        assert blue_colors != gold_colors


# ── page size / margins ───────────────────────────────────────────────────────


class TestPage:
    def test_a4_page(self):
        doc = _open(render_layout_docx(_comprehensive_cvm(), sidebar_layout(), blue_theme()))
        section = doc.sections[0]
        assert abs(section.page_width.mm - 210.0) < 1
        assert abs(section.page_height.mm - 297.0) < 1
        assert section.left_margin.mm > 0 and section.top_margin.mm > 0

    def test_letter_page(self):
        cvm = ContentView(
            stable_id="resume.letter",
            profile=Profile(full_name="Jane Doe", professional_title="Engineer"),
        )
        # Build a tree with a Letter page by overriding the layout page size.
        from app.rendering.builder import TreeBuilder
        from app.rendering.context import RenderContext, RenderState
        from app.rendering.layout.layout_definition import PageOptions
        from app.rendering.layout_html import default_component_registry

        layout = sidebar_layout().model_copy(
            update={"page": PageOptions(page_size="Letter")}
        )
        context = RenderContext(layout=layout, theme=blue_theme(), state=RenderState())
        tree = TreeBuilder(default_component_registry()).build(cvm, context)
        doc = _open(RenderTreeDOCXRenderer().render(tree, theme=blue_theme()))
        section = doc.sections[0]
        assert abs(section.page_width.mm - 215.9) < 1
        assert abs(section.page_height.mm - 279.4) < 1


# ── multi-page long resume ────────────────────────────────────────────────────


class TestLongResume:
    def test_long_content_all_present(self):
        doc = _open(render_layout_docx(_long_cvm(), sidebar_layout(), blue_theme()))
        text = _doc_text(doc)
        assert "Company 5 International Solutions Group" in text  # 6th job
        assert "Certification 3" in text
        assert "Project 2" in text
        assert "Award 1" in text
        assert "Language 2" in text
        # Long document spans many paragraphs (body + table cells).
        total_paragraphs = len(doc.paragraphs) + sum(
            len(cell.paragraphs)
            for table in doc.tables
            for row in table.rows
            for cell in row.cells
        )
        assert total_paragraphs > 50


# ── hyperlinks ────────────────────────────────────────────────────────────────


class TestHyperlinks:
    def test_clickable_hyperlinks_and_targets(self):
        doc = _open(render_layout_docx(_comprehensive_cvm(), sidebar_layout(), blue_theme()))
        targets = [
            rel.target_ref
            for rel in doc.part.rels.values()
            if "hyperlink" in rel.reltype
        ]
        assert "https://linkedin.com/in/sharma" in targets
        assert "https://github.com/sharma" in targets
        assert "https://sharma.dev" in targets
        assert "https://github.com/example/alpha" in targets

    def test_link_text_present(self):
        text = _doc_text(_open(render_layout_docx(_comprehensive_cvm(), sidebar_layout(), blue_theme())))
        assert "https://linkedin.com/in/sharma" in text


# ── special characters ────────────────────────────────────────────────────────


class TestSpecialCharacters:
    def test_special_chars_survive_as_text(self):
        cvm = ContentView(
            stable_id="resume.special",
            profile=Profile(full_name="R&D C++ / C# Lead", professional_title="Senior Engineer — Platform"),
            summary="Value & cost analysis <50% + #priority",
        )
        text = _doc_text(_open(render_layout_docx(cvm, classic_layout(), blue_theme())))
        assert "R&D C++ / C# Lead" in text
        assert "Senior Engineer — Platform" in text
        assert "Value & cost analysis" in text


# ── determinism ───────────────────────────────────────────────────────────────


class TestDeterminism:
    def test_structured_content_deterministic(self):
        cvm = _comprehensive_cvm()
        a = _doc_text(_open(render_layout_docx(cvm, sidebar_layout(), blue_theme())))
        b = _doc_text(_open(render_layout_docx(cvm, sidebar_layout(), blue_theme())))
        assert a == b


# ── immutability ──────────────────────────────────────────────────────────────


class TestImmutability:
    def test_cvm_and_tree_not_mutated(self):
        cvm = _comprehensive_cvm()
        hash_before = cvm.content_hash
        from app.rendering.builder import TreeBuilder
        from app.rendering.context import RenderContext, RenderState
        from app.rendering.layout_html import default_component_registry

        context = RenderContext(layout=sidebar_layout(), theme=blue_theme(), state=RenderState())
        tree = TreeBuilder(default_component_registry()).build(cvm, context)
        tree_before = tree.model_dump_json()
        RenderTreeDOCXRenderer().render(tree, theme=blue_theme())
        assert cvm.content_hash == hash_before
        assert tree.model_dump_json() == tree_before


# ── error handling ────────────────────────────────────────────────────────────


class TestErrors:
    def test_invalid_tree_rejected(self):
        from app.rendering.tree import A4, PageMargins, TextData

        text = RenderNode(id="t1", kind=NodeKind.TEXT, data=TextData(type="text", text="x"))
        block = RenderNode(id="b1", kind=NodeKind.BLOCK, region="main", children=(text,))
        section = RenderNode(id="s1", kind=NodeKind.SECTION, content_ref="summary", region="main", children=(block,))
        region = RenderNode(id="r1", kind=NodeKind.REGION, region="main", children=(section,))
        page = RenderNode(id="p1", kind=NodeKind.PAGE, page_size=A4, margins=PageMargins(), children=(region,))
        invalid = RenderNode(id="doc", kind=NodeKind.DOCUMENT, children=(page,))
        with pytest.raises(TreeValidationError):
            RenderTreeDOCXRenderer().render(invalid)

    def test_empty_document_renders(self):
        tree = RenderNode(id="doc", kind=NodeKind.DOCUMENT, children=())
        data = RenderTreeDOCXRenderer().render(tree)
        assert data[:2] == b"PK"


# ── renderer import isolation ─────────────────────────────────────────────────


def _module_imports(dotted: str) -> set[str]:
    module = importlib.import_module(dotted)
    path = Path(inspect.getsourcefile(module))
    tree = ast.parse(path.read_text(encoding="utf-8"))
    imports: set[str] = set()
    for node in ast.walk(tree):
        if isinstance(node, ast.ImportFrom) and node.module:
            imports.add(node.module)
        elif isinstance(node, ast.Import):
            for alias in node.names:
                imports.add(alias.name.split(".")[0])
    return imports


def _imports_package(imports: set[str], package: str) -> bool:
    prefix = package + "."
    return any(item == package or item.startswith(prefix) for item in imports)


class TestImportIsolation:
    def test_docx_renderer_is_pure_render_tree_consumer(self):
        imports = _module_imports("app.rendering.renderers.tree_docx_renderer")
        assert _imports_package(imports, "app.rendering.tree")
        assert _imports_package(imports, "app.rendering.theme")
        assert _imports_package(imports, "app.rendering.common")
        for forbidden in (
            "app.rendering.content",
            "app.rendering.components",
            "app.rendering.layout",
            "app.rendering.preview",
            "app.rendering.registry",
            "app.rendering.builder",
            "app.services",
            "app.models",
        ):
            assert not _imports_package(imports, forbidden), forbidden
